"""
Charts for the Revenue Model Explainer, read only from the calculated workbook's
Summary sheet, so a figure cannot disagree with the table printed beside it.
Tokenomics.net palette, styled as the Phase 3 and Phase 5 charts.

Run: python charts_rev.py      -> ../charts/*.png
"""

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.ticker import FuncFormatter

GOLD = "#B8956E"; DARK_GOLD = "#96734A"; LIGHT_GOLD = "#D4B896"
DARK = "#1A1714"; WARM = "#D4CFC8"; MED = "#9A9590"; CREAM = "#FAF8F5"
RED = "#A94442"; GREEN = "#4A7C59"

HERE = Path(__file__).parent
PHASE = HERE.parent
WB = PHASE / "tools" / "Aurumix_Revenue_Model_calculated.xlsx"
CH = PHASE / "charts"
CH.mkdir(exist_ok=True)

ws = openpyxl.load_workbook(WB, data_only=True)["Summary"]
ROWS = {}
for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=10, values_only=True):
    if isinstance(r[0], str) and r[0].strip() not in ROWS:
        ROWS[r[0].strip()] = [float(v or 0) for v in r[3:10]]
YEARS = ["Y1", "Y2", "Y3", "Y4", "Y5", "Y6", "Y7"]
X = np.arange(7)


def row(label):
    return np.array(ROWS[label])


def usd_m(v, _=None):
    if v == 0:
        return "0"
    sign = "-" if v < 0 else ""
    v = abs(v)
    return sign + ("USD %.1fm" % (v / 1e6) if v >= 1e6 else "USD %.0fk" % (v / 1e3))


def style(ax, fig):
    ax.set_facecolor(CREAM)
    fig.patch.set_facecolor(CREAM)
    ax.grid(True, axis="y", alpha=0.18, color=MED, linewidth=0.7)
    ax.set_axisbelow(True)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(WARM)
    ax.tick_params(colors=DARK, labelsize=9)


def save(fig, name):
    fig.tight_layout()
    fig.savefig(CH / name, dpi=160, facecolor=CREAM)
    plt.close(fig)
    print("  wrote", name)


# 1. revenue by stream -------------------------------------------------------
streams = [
    ("Entry fee (SIP + spot)", row("Stream 1a: Entry fee - SIP") + row("Stream 1b: Entry fee - SPOT"), GOLD),
    ("B2B platform fee", row("Stream 6: B2B platform fee"), DARK),
    ("Cardholder fees", row("Stream 4: Cardholder fees"), DARK_GOLD),
    ("Family plan and Digital Will", row("Stream 3: Family plan and Digital Will"), LIGHT_GOLD),
    ("Interchange and lending", row("Stream 2: Card interchange") + row("Stream 5: Lending revenue share"), MED),
]
fig, ax = plt.subplots(figsize=(8, 4.2))
style(ax, fig)
base = np.zeros(7)
for lab, v, c in streams:
    ax.bar(X, v, 0.62, bottom=base, color=c, label=lab)
    base += v
total = row("TOTAL REVENUE")
assert np.allclose(base, total), "streams do not reconcile to total revenue"
for i, t in enumerate(total):
    ax.text(i, t + total.max() * 0.012, usd_m(t), ha="center", va="bottom", fontsize=8, color=DARK)
ax.set_xticks(X, YEARS)
ax.yaxis.set_major_formatter(FuncFormatter(usd_m))
ax.set_ylim(0, total.max() * 1.1)
h, l = ax.get_legend_handles_labels()
ax.legend(h[::-1], l[::-1], fontsize=8, frameon=False, loc="upper left")
save(fig, "rev_by_stream.png")

# 2. customers ---------------------------------------------------------------
pay, hold, cards = row("Paying customers"), row("Holders (stopped paying, still hold gold)"), row("Active cards")
fig, ax = plt.subplots(figsize=(8, 3.8))
style(ax, fig)
ax.bar(X, pay, 0.62, color=GOLD, label="Paying customers")
ax.bar(X, hold, 0.62, bottom=pay, color=LIGHT_GOLD, label="Holders (stopped paying, still hold gold)")
ax.plot(X, cards, color=DARK, marker="o", ms=4, lw=1.8, label="Active cards")
assert np.allclose(pay + hold, row("Cumulative ever acquired")), "paying + holders != ever acquired"
for i in range(7):
    ax.text(i, pay[i] + hold[i] + 2500, "{:,.0f}".format(pay[i] + hold[i]), ha="center", fontsize=8, color=DARK)
ax.set_xticks(X, YEARS)
ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: "{:,.0f}k".format(v / 1e3) if v else "0"))
ax.set_ylim(0, (pay + hold).max() * 1.12)
ax.legend(fontsize=8, frameon=False, loc="upper left")
save(fig, "customers.png")

# 3. profit ------------------------------------------------------------------
cost = row("TOTAL COST BASE")
net = row("NET PROFIT")
cum = row("CUMULATIVE NET PROFIT")
assert np.allclose(total - cost, net), "revenue - cost does not equal net profit"
fig, ax = plt.subplots(figsize=(8, 4))
style(ax, fig)
w = 0.36
ax.bar(X - w / 2, total, w, color=GOLD, label="Total revenue")
ax.bar(X + w / 2, cost, w, color=MED, label="Total cost base")
ax.plot(X, cum, color=DARK, marker="o", ms=4, lw=1.8, label="Cumulative net profit")
ax.axhline(0, color=DARK, lw=0.8)
for i, v in enumerate(net):
    ax.text(i, max(total[i], cost[i]) + 90000, usd_m(v), ha="center", fontsize=7.5,
            color=GREEN if v >= 0 else RED)
ax.set_xticks(X, YEARS)
ax.yaxis.set_major_formatter(FuncFormatter(usd_m))
ax.set_ylim(min(cum.min(), 0) * 1.15, max(total.max(), cost.max()) * 1.12)
ax.legend(fontsize=8, frameon=False, loc="upper left")
ax.text(0.99, 0.02, "Figures above the bars: net profit that year (upper bound, before headcount and tax)",
        transform=ax.transAxes, ha="right", va="bottom", fontsize=7.5, color=MED)
save(fig, "profit.png")

print("revenue Y7 %.0f | net Y7 %.0f | cum Y7 %.0f" % (total[-1], net[-1], cum[-1]))
